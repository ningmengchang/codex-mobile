#!/usr/bin/env python3
import argparse
import datetime
import json
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_ROWS = 200
MAX_COLUMNS = 50


def color_value(color):
    if color is None or color.type != "rgb" or not color.rgb:
        return None
    value = str(color.rgb)[-6:]
    return f"#{value}" if value.upper() != "000000" else "#000000"


def display_value(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if "%" in (cell.number_format or ""):
            return f"{value * 100:g}%"
        return f"{value:g}"
    return str(value)


def cell_style(cell):
    fill = cell.fill
    font = cell.font
    alignment = cell.alignment
    border = cell.border
    background = color_value(fill.fgColor) if fill and fill.fill_type else None
    return {
        "background": background,
        "color": color_value(font.color) if font else None,
        "bold": bool(font and font.bold),
        "italic": bool(font and font.italic),
        "horizontal": alignment.horizontal if alignment else None,
        "vertical": alignment.vertical if alignment else None,
        "wrap": bool(alignment and alignment.wrap_text),
        "border": any(getattr(border, side).style for side in ("left", "right", "top", "bottom")),
    }


def workbook_metadata(workbook):
    sheets = []
    for index, sheet in enumerate(workbook.worksheets):
        sheets.append({
            "index": index,
            "name": sheet.title,
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "hidden": sheet.sheet_state != "visible",
        })
    return {"sheets": sheets, "maxRows": MAX_ROWS, "maxColumns": MAX_COLUMNS}


def worksheet_data(workbook, index):
    if index < 0 or index >= len(workbook.worksheets):
        raise ValueError("工作表不存在")
    sheet = workbook.worksheets[index]
    rendered_rows = min(sheet.max_row, MAX_ROWS)
    rendered_columns = min(sheet.max_column, MAX_COLUMNS)
    styles = [{}]
    style_indexes = {json.dumps({}, sort_keys=True): 0}
    rows = []
    for row_number in range(1, rendered_rows + 1):
        row = []
        for column_number in range(1, rendered_columns + 1):
            cell = sheet.cell(row=row_number, column=column_number)
            style = cell_style(cell) if cell.has_style else {}
            key = json.dumps(style, sort_keys=True)
            if key not in style_indexes:
                style_indexes[key] = len(styles)
                styles.append(style)
            row.append([display_value(cell), style_indexes[key]])
        rows.append(row)
    merges = []
    for merged in sheet.merged_cells.ranges:
        if merged.min_row > rendered_rows or merged.min_col > rendered_columns:
            continue
        merges.append({
            "startRow": merged.min_row,
            "startColumn": merged.min_col,
            "endRow": min(merged.max_row, rendered_rows),
            "endColumn": min(merged.max_col, rendered_columns),
        })
    widths = []
    for column_number in range(1, rendered_columns + 1):
        width = sheet.column_dimensions[get_column_letter(column_number)].width
        widths.append(max(48, min(420, round((width or 9) * 7 + 12))))
    heights = []
    for row_number in range(1, rendered_rows + 1):
        height = sheet.row_dimensions[row_number].height
        heights.append(max(24, min(180, round((height or 18) * 1.34))))
    return {
        "index": index,
        "name": sheet.title,
        "rows": sheet.max_row,
        "columns": sheet.max_column,
        "renderedRows": rendered_rows,
        "renderedColumns": rendered_columns,
        "truncated": sheet.max_row > rendered_rows or sheet.max_column > rendered_columns,
        "styles": styles,
        "data": rows,
        "merges": merges,
        "widths": widths,
        "heights": heights,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=("workbook", "sheet"))
    parser.add_argument("path")
    parser.add_argument("index", nargs="?", type=int)
    args = parser.parse_args()
    workbook = load_workbook(args.path, read_only=False, data_only=False, keep_links=False)
    try:
        result = workbook_metadata(workbook) if args.command == "workbook" else worksheet_data(workbook, args.index)
        json.dump(result, sys.stdout, ensure_ascii=False, separators=(",", ":"))
    finally:
        workbook.close()


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
